"""
Opens the original DOCX and replaces only the target paragraph texts
while preserving all paragraph-level styles, bullets, and spacing.
"""

import copy
import shutil
import os
from docx import Document
from docx.oxml.ns import qn


def _replace_paragraph_text(para, new_text: str):
    """
    Replace the text of a paragraph while preserving the run formatting of the first run.
    All existing runs are cleared and replaced with a single run carrying the first run's rPr.
    """
    if not para.runs:
        # No runs — just set via XML directly
        for child in list(para._p):
            if child.tag.endswith("}r"):
                para._p.remove(child)
        run = para.add_run(new_text)
        return

    # Capture formatting from the first run
    first_run = para.runs[0]
    first_rpr = copy.deepcopy(first_run._r.find(qn("w:rPr")))

    # Remove all existing runs
    for run in para.runs:
        para._p.remove(run._r)

    # Add a new single run with the preserved formatting
    new_run = para.add_run(new_text)
    if first_rpr is not None:
        existing_rpr = new_run._r.find(qn("w:rPr"))
        if existing_rpr is not None:
            new_run._r.remove(existing_rpr)
        new_run._r.insert(0, first_rpr)


def update_docx(
    source_path: str,
    output_path: str,
    summary_updates: dict | None,
    bullet_updates: list[dict] | None,
) -> str:
    """
    Create a tailored copy of the resume DOCX.

    Args:
        source_path: path to the original DOCX
        output_path: where to write the tailored DOCX
        summary_updates: {"paragraph_indexes": [int], "rewritten_text": str}
                         OR None to skip summary update
        bullet_updates: [{"paragraph_index": int, "rewritten_text": str}, ...]
                         OR None to skip bullet updates

    Returns:
        output_path on success
    """
    # Work on a copy so the original is never modified
    shutil.copy2(source_path, output_path)
    doc = Document(output_path)
    paragraphs = doc.paragraphs

    # --- Summary ---
    if summary_updates:
        indexes = summary_updates.get("paragraph_indexes", [])
        new_text = summary_updates.get("rewritten_text", "")
        if indexes and new_text:
            # Write the full rewritten summary into the first paragraph index,
            # clear the rest (keeps structure intact)
            _replace_paragraph_text(paragraphs[indexes[0]], new_text)
            for idx in indexes[1:]:
                _replace_paragraph_text(paragraphs[idx], "")

    # --- Bullets ---
    if bullet_updates:
        for update in bullet_updates:
            idx = update.get("paragraph_index")
            new_text = update.get("rewritten_text", "")
            if idx is not None and new_text and idx < len(paragraphs):
                _replace_paragraph_text(paragraphs[idx], new_text)

    doc.save(output_path)
    return output_path
